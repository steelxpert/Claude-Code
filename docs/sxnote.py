"""Shared page furniture for English STEELXPERT technical notes (wraps the house `common.py` helper).

Usage in a build script:
    N = sxnote.Note(common_path, signature_path, projno, docid, project, date_txt, date_num, date_tup, rev)
    N.H(1, "Heading", new_page=True); N.P("text with **bold** and x_[sub]"); N.table(...); N.figure(...); N.signature(); N.save(out)
"""
from __future__ import annotations

import importlib.util, zipfile, shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FIG_W_CM, FIG_MIN_CM = 15.5, 14.2


class Note:
    def __init__(self, common_path, signature_path, projno, docid, project, date_txt, date_num, date_tup, rev="rev0",
                 title="", subject="", keywords="", kind="Technical note", final=False):
        spec = importlib.util.spec_from_file_location("common", common_path)
        C = importlib.util.module_from_spec(spec); spec.loader.exec_module(C)
        C.LANG = "en-GB"; C.HIGHLIGHT[0] = not final
        self.C, self.sig, self.projno, self.docid, self.project, self.rev = C, signature_path, projno, docid, project, rev
        self.date_txt, self.date_num = date_txt, date_num
        self.nfig = self.ntab = 0
        doc = self.doc = Document(); C.DOCREF[0] = doc
        C.set_lang(doc); C.base_styles(doc); C.page_setup(doc.sections[0])
        C.core_props(doc, title, subject, keywords, date_tup); doc.core_properties.category = kind
        self.kind = kind
        self._header(doc.sections[0].header); self._footer(doc.sections[0].footer)

    # ---- furniture
    def _header(self, hdr):
        C = self.C
        p = hdr.paragraphs[0]; C.rich(p, self.kind, size=11, bold=True, color="404040"); p.paragraph_format.space_after = Pt(3)
        t = hdr.add_table(rows=3, cols=4, width=Cm(C.CONTENT_W)); C.set_widths(t, [2.4, 10.1, 2.0, 2.0])
        vals = [("Project:", self.project, "Page:", None), ("Project No.:", self.projno, "Revision:", self.rev), ("Document:", self.docid, "", "")]
        for i, row in enumerate(vals):
            for j, v in enumerate(row):
                cl = t.cell(i, j)
                if v is None:
                    pp = cl.paragraphs[0]; pp.paragraph_format.space_after = Pt(0); pp.paragraph_format.line_spacing = 1.0
                    pp.style = C.small_style(name="Header small"); C.field(pp, "PAGE", C.HDR_SIZE)
                else:
                    C.cell_text(cl, v, size=C.HDR_SIZE)
                C.borders(cl, bottom=6 if i == 2 else None)
        C.cell_margins(t, top=45, bottom=15, left=60, right=60)
        for cl in t.rows[2].cells:
            tcPr = cl._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar"); b = OxmlElement("w:bottom"); b.set(qn("w:w"), "90"); b.set(qn("w:type"), "dxa"); m.append(b); tcPr.append(m)

    def _footer(self, ftr):
        C = self.C
        p = ftr.paragraphs[0]; p.style = C.small_style(name="Footer small", color="404040"); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.tab_stops.add_tab_stop(Cm(C.CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
        C.rich(p, f"STEELXPERT Kft. · Csaba Bán, MSc civil engineer (T-T/01-18387) · {self.date_num} · {self.rev}\tPage: ")
        C.field(p, "PAGE", C.HDR_SIZE); C.rich(p, " / "); C.field(p, "NUMPAGES", C.HDR_SIZE)

    def cover(self, title, subtitle, rows, rev_rows):
        C, doc = self.C, self.doc
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(60); C.rich(p, self.kind.upper(), size=10, bold=True, color="808080")
        p = doc.add_paragraph(); C.rich(p, title, size=20, bold=True); p.paragraph_format.space_after = Pt(4)
        p = doc.add_paragraph(); C.rich(p, subtitle, size=13); p.paragraph_format.space_after = Pt(24)
        C.info_table(doc, rows); doc.add_paragraph().paragraph_format.space_after = Pt(12)
        self.table("Revision history", ["Rev.", "Description", "Date", "Prepared", "Checked"], rev_rows, [1.3, 8.2, 2.4, 2.3, 2.3], letter="A")

    def toc(self):
        self.H(1, "Contents", new_page=True)
        p = self.doc.add_paragraph()
        def fld(t): r = p.add_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e)
        fld("begin"); r = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = ' TOC \\o "1-3" \\h \\z \\u '; r._r.append(it)
        fld("separate"); p.add_run("Table of contents: right-click and update the field (F9) in Word.").italic = True; fld("end")

    # ---- content
    def H(self, level, text, new_page=False):
        h = self.doc.add_heading(text, level=level)
        if new_page: h.paragraph_format.page_break_before = True
        return h

    def P(self, text, **kw):
        p = self.doc.add_paragraph(); self.C.rich(p, text, **kw); return p

    def B(self, text):
        p = self.doc.add_paragraph(style="List Bullet"); self.C.rich(p, text); return p

    def N(self, items):
        for i, r in enumerate(items, 1):
            p = self.doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.8)); self.C.rich(p, f"{i}.\t{r}")

    def table(self, caption, header, rows, widths, letter=None, size=9):
        C, doc = self.C, self.doc
        if letter is None: self.ntab += 1; no = str(self.ntab)
        else: no = letter
        cap = doc.add_paragraph(style="Caption"); C.rich(cap, f"Table {no}: {caption}"); cap.paragraph_format.keep_with_next = True
        t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.style = doc.styles["Table Grid"]; C.set_widths(t, widths)
        C.cell_margins(t, top=45, bottom=30, left=60, right=60)
        for j, h in enumerate(header): C.cell_text(t.cell(0, j), h, size=size, bold=True); C.shade(t.cell(0, j), "E6E6E6")
        C.cant_split(t.rows[0], header=True)
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r): C.cell_text(t.cell(i, j), str(v), size=size)
            C.cant_split(t.rows[i])
        C.keep_table_start(t); doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return no

    def figure(self, path, caption, width_cm=FIG_W_CM):
        assert width_cm >= FIG_MIN_CM, f"figure narrower than {FIG_MIN_CM} cm"
        self.nfig += 1
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        cap = self.doc.add_paragraph(style="Caption"); self.C.rich(cap, f"Figure {self.nfig}: {caption}")
        return self.nfig

    def signature(self):
        C, doc = self.C, self.doc
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18); p.paragraph_format.keep_with_next = True
        C.rich(p, f"Budapest, {self.date_txt}", size=10)
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(self.sig, width=Cm(4.0))
        lines = ["- - - - - - - - - - - - - - - - - - - - - - -", "**STEELXPERT Kft.**", "Csaba Bán", "MSc civil engineer", "lead structural designer", "T-T/01-18387"]
        for i, ln in enumerate(lines):
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i < len(lines) - 1: p.paragraph_format.keep_with_next = True
            C.rich(p, ln, size=10)

    def save(self, out):
        out = str(out); Path(out).parent.mkdir(parents=True, exist_ok=True); self.doc.save(out)
        tmp = out + ".tmp"
        with zipfile.ZipFile(out) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.endswith(".xml"):
                    data = data.replace(b'w:val="en-US"', b'w:val="en-GB"').replace(b'w:eastAsia="en-US"', b'w:eastAsia="en-GB"')
                zout.writestr(item, data)
        shutil.move(tmp, out)
        return sorted(set(self.C.PENDING))
